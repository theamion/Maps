#!/usr/bin/env python3
"""
apply_osm_coords.py - copy 'OSM lat' / 'OSM lon' over 'Latitude' /
'Longitude' in the Junctions tab of junctions_topology_v5.xlsx, for every
junction where both OSM values are filled in, and save the workbook.

Usage:
    python3 apply_osm_coords.py
    python3 apply_osm_coords.py --xlsx other.xlsx
"""
import argparse
import os

from openpyxl import load_workbook

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_XLSX = os.path.join(SCRIPT_DIR, "junctions_topology_v5.xlsx")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    args = ap.parse_args()

    wb = load_workbook(args.xlsx)
    ws = wb["Junctions"]
    col = {c.value: c.column for c in ws[1] if c.value}
    for h in ("Latitude", "Longitude", "OSM lat", "OSM lon"):
        if h not in col:
            raise SystemExit(f"Column '{h}' not found in the Junctions tab of {args.xlsx}")

    changed = 0
    for r in range(2, ws.max_row + 1):
        osm_lat = ws.cell(r, col["OSM lat"]).value
        osm_lon = ws.cell(r, col["OSM lon"]).value
        if osm_lat is None or osm_lon is None:
            continue
        ws.cell(r, col["Latitude"]).value = osm_lat
        ws.cell(r, col["Longitude"]).value = osm_lon
        changed += 1

    wb.save(args.xlsx)
    print(f"{changed} junctions updated with OSM coordinates, saved to {args.xlsx}")


if __name__ == "__main__":
    main()
