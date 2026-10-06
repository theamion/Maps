#!/usr/bin/env python3
"""
serways_brands.py — Find the fast-food/coffee brand(s) at each German
Autobahn fuel station in the Points tab, from serways.de (the consumer
site of Tank & Rast, which operates almost every Autobahn service area),
and write them into a 'Food brand(s)' column mapmaking.py reads.

HOW A STATION IS MATCHED
-------------------------
serways.de publishes one page per service area, named like
'brohltal-ost' or 'kamener-kreuz' (no direction suffix when there's only
one side). A Points row is matched by normalising its Name the same way
(umlauts expanded, accents stripped, lower-cased) and comparing it to
that slug, with or without a trailing east/west/north/south. A name
that already ends in a direction (e.g. 'Denkendorf Nord') is matched to
that exact side; a plain name (e.g. 'Brohltal') that has both an Ost and
a West page on serways.de is matched to both — the two sides are kept
separate throughout, since they can be (and often are) a different food
brand, or even a different fuel brand.

Only a minority of Points rows match: most are Dutch/Belgian/Austrian
stations outside the Tank & Rast network, and some German ones use a
different name than serways.de's (the run prints these so you can look
them up and add a line to NAME_OVERRIDES below, or just tell Claude).

WHAT GETS WRITTEN
-------------------
'Food brand(s)' in the Points tab, written only for a matched row — any
row Code can't match is left exactly as it was, nothing is cleared. The
column uses the same free-text convention as 'Fuel brand' already does:
'Dallmayr, McDonald's' (same on both sides, or only one side exists) or
'McDonald's east; Dallmayr west' (a genuinely different side pages). It
never touches 'Fuel brand' or 'Facilities'.

RATE LIMITING / RESUMABILITY
-----------------------------
Same approach as the osm_*.py scripts: a small delay between requests and
a JSON cache keyed by serways.de URL slug, so an interrupted run can be
resumed by just running the script again - already-cached pages aren't
re-fetched.

Usage:
    python3 serways_brands.py
    python3 serways_brands.py --limit 10 --output test.xlsx
"""
import argparse
import json
import os
import re
import time
import unicodedata
import urllib.request

from openpyxl import load_workbook

import map_icons

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_XLSX = os.path.join(SCRIPT_DIR, "junctions_topology_v5.xlsx")
DEFAULT_CACHE = os.path.join(SCRIPT_DIR, "serways_cache.json")
SITEMAP_URL = "https://www.serways.de/standort-sitemap.xml"
DIRS = ("ost", "west", "nord", "sued")
DIR_TO_COMPASS = {"ost": "east", "west": "west", "nord": "north", "sued": "south"}

KNOWN_FUEL = [name for name, *_ in map_icons.FUEL_BRAND_COLOURS]
KNOWN_FOOD = [name for name, *_ in map_icons.FOOD_BRAND_COLOURS]

# Points 'Name' -> the serways.de base slug (before any -ost/-west/...),
# for a station whose name doesn't normalise to a close enough match on
# its own. Add a line here (or just tell Claude the right slug) for a
# German station the run lists as unmatched.
NAME_OVERRIDES = {}


def norm(s):
    s = s.lower()
    s = s.replace('ä', 'ae').replace('ö', 'oe').replace('ü', 'ue').replace('ß', 'ss')
    s = unicodedata.normalize('NFKD', s)
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return re.sub(r'[^a-z0-9]+', '-', s).strip('-')


def simplify(key):
    return key.replace('ae', 'a').replace('oe', 'o').replace('ue', 'u')


def fetch(url, opener, timeout=20):
    req = urllib.request.Request(url, headers={
        'User-Agent': 'Mozilla/5.0 (compatible; personal route-map research; '
                      'github.com/theamion/Maps)'})
    return opener.open(req, timeout=timeout).read().decode('utf-8', 'ignore')


def load_slug_index(opener):
    xml = fetch(SITEMAP_URL, opener)
    slugs = sorted(set(u.rstrip('/').rsplit('/', 1)[-1]
                        for u in re.findall(r'<loc>([^<]+)</loc>', xml)))
    base_to_variants = {}
    for sn in slugs:
        parts = sn.rsplit('-', 1)
        base, d = (parts[0], parts[1]) if len(parts) == 2 and parts[1] in DIRS else (sn, None)
        base_to_variants.setdefault(simplify(base), []).append((d, sn))
    return base_to_variants


def match_slugs(name, base_to_variants):
    if name in NAME_OVERRIDES:
        base = simplify(norm(NAME_OVERRIDES[name]))
        return [sn for _d, sn in base_to_variants.get(base, [])]
    key = simplify(norm(name))
    m = re.match(r'^(.*)-(ost|west|nord|sued)$', key)
    base_part, dir_part = (m.group(1), m.group(2)) if m else (key, None)
    cands = base_to_variants.get(base_part)
    if not cands:
        return []
    if dir_part:
        exact = [sn for d, sn in cands if d == dir_part]
        if exact:
            return exact
    return [sn for _d, sn in cands]


def brands_on_page(html):
    alts = re.findall(r'<img[^>]*alt="([^"]+)"', html)
    fuel = [a for a in KNOWN_FUEL if any(a.lower() == x.lower() for x in alts)]
    food = [a for a in KNOWN_FOOD if any(a.lower() == x.lower() for x in alts)]
    return fuel, food


def get_slug_brands(slug, cache, opener, delay):
    if slug in cache:
        return cache[slug]
    url = f"https://www.serways.de/standorte/{slug}/"
    try:
        html = fetch(url, opener)
        fuel, food = brands_on_page(html)
        cache[slug] = {"status": "OK", "fuel": fuel, "food": food}
    except Exception as e:
        cache[slug] = {"status": f"ERROR: {e}", "fuel": [], "food": []}
    time.sleep(delay)
    return cache[slug]


def format_brand_text(by_slug_dir, field):
    """by_slug_dir: {direction-or-None: {'fuel':[...], 'food':[...]}} for
    the 1 or 2 slugs matched for one point. Builds the 'X east; Y west'
    (or plain 'X, Y') text for `field` ('fuel' or 'food')."""
    sides = {d: v[field] for d, v in by_slug_dir.items()}
    if len(sides) == 1:
        brands = next(iter(sides.values()))
        return ', '.join(brands) if brands else None
    values = list(sides.values())
    if all(set(v) == set(values[0]) for v in values):
        brands = values[0]
        return ', '.join(brands) if brands else None
    clauses = []
    for d, brands in sides.items():
        if not brands:
            continue
        compass = DIR_TO_COMPASS.get(d, d)
        clauses.append(f"{', '.join(brands)} {compass}")
    return '; '.join(clauses) if clauses else None


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    ap.add_argument("--output", default=None, help="write to a copy instead of overwriting --xlsx")
    ap.add_argument("--cache", default=DEFAULT_CACHE)
    ap.add_argument("--delay", type=float, default=0.4, help="seconds between page fetches")
    ap.add_argument("--limit", type=int, default=None, help="only process the first N matched points (for testing)")
    args = ap.parse_args()

    opener = urllib.request.build_opener()
    cache = json.load(open(args.cache, encoding='utf-8')) if os.path.exists(args.cache) else {}

    wb = load_workbook(args.xlsx)
    ws = wb["Points"]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header)}
    for col in ("Point ID", "Name", "Category"):
        if col not in idx:
            raise SystemExit(f"Column '{col}' not found in the Points tab of {args.xlsx}")

    print("Fetching the serways.de station list...")
    base_to_variants = load_slug_index(opener)

    rows = []
    for r in range(2, ws.max_row + 1):
        pid = ws.cell(r, idx["Point ID"] + 1).value
        cat = ws.cell(r, idx["Category"] + 1).value
        name = ws.cell(r, idx["Name"] + 1).value
        if pid and cat in ("tankstation", "Autohof"):
            rows.append((r, name))

    matched = [(r, name, match_slugs(name, base_to_variants)) for r, name in rows]
    matched_rows = [(r, name, slugs) for r, name, slugs in matched if slugs]
    unmatched = [name for _r, name, slugs in matched if not slugs]
    if args.limit:
        matched_rows = matched_rows[:args.limit]

    print(f"{len(rows)} tankstation/Autohof points, {len(matched_rows)} matched to a serways.de page"
          f" ({len(matched)-len(matched_rows)} not matched)")

    all_slugs = sorted({s for _r, _n, slugs in matched_rows for s in slugs})
    print(f"Fetching {len(all_slugs)} station pages (cached ones are instant)...")
    for i, slug in enumerate(all_slugs):
        get_slug_brands(slug, cache, opener, args.delay)
        if (i + 1) % 25 == 0:
            print(f"  ...{i+1}/{len(all_slugs)}")
            json.dump(cache, open(args.cache, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)
    json.dump(cache, open(args.cache, 'w', encoding='utf-8'), ensure_ascii=False, indent=1)

    existing = {c.value: c.column for c in ws[1] if c.value}
    col_name = "Food brand(s)"
    if col_name not in existing:
        col = ws.max_column + 1
        ws.cell(row=1, column=col, value=col_name)
        existing[col_name] = col
    food_col = existing[col_name]

    updated = 0
    for r, name, slugs in matched_rows:
        by_dir = {}
        for slug in slugs:
            parts = slug.rsplit('-', 1)
            d = parts[1] if len(parts) == 2 and parts[1] in DIRS else None
            by_dir[d] = get_slug_brands(slug, cache, opener, 0)
        food_text = format_brand_text(by_dir, "food")
        if food_text:
            ws.cell(row=r, column=food_col, value=food_text)
            updated += 1

    out_path = args.output or args.xlsx
    wb.save(out_path)
    print(f"Wrote '{col_name}' for {updated} points -> {out_path}")

    if unmatched:
        print(f"\n{len(unmatched)} tankstation/Autohof points had no matching serways.de page")
        print("(expected for non-German stations; a German name here is worth a manual look):")
        for name in unmatched:
            print(" ", name)


if __name__ == "__main__":
    main()
