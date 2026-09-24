#!/usr/bin/env python3
"""
apply_osm_distances.py - copy 'OSRM route afstand (km)' over 'Distance (km)'
in the Segments tab of junctions_topology_v5.xlsx, for every segment whose
'OSRM status' is OK, and save the workbook. 'Distance (km)' is what
mapmaking.py uses for its distance labels and route lengths.

Usage:
    python3 apply_osm_distances.py
    python3 apply_osm_distances.py --xlsx other.xlsx
"""
import argparse
import os

from openpyxl import load_workbook

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
DEFAULT_XLSX = os.path.join(SCRIPT_DIR, "junctions_topology_v5.xlsx")


def main():
    ap = argparse.ArgumentParser(description=__doc__, formatter_class=argparse.RawDescriptionHelpFormatter)
    ap.add_argument("--xlsx", default=DEFAULT_XLSX)
    args = ap.parse_args()

    wb = load_workbook(args.xlsx)
    ws = wb["Segments"]
    col = {c.value: c.column for c in ws[1] if c.value}
    for h in ("Distance (km)", "OSRM status", "OSRM route afstand (km)"):
        if h not in col:
            raise SystemExit(f"Column '{h}' not found in the Segments tab of {args.xlsx}")

    changed = 0
    for r in range(2, ws.max_row + 1):
        if ws.cell(r, col["OSRM status"]).value != "OK":
            continue
        osrm_km = ws.cell(r, col["OSRM route afstand (km)"]).value
        if osrm_km is None:
            continue
        ws.cell(r, col["Distance (km)"]).value = osrm_km
        changed += 1

    wb.save(args.xlsx)
    print(f"{changed} segments updated with OSRM distances, saved to {args.xlsx}")


if __name__ == "__main__":
    main()
