#!/usr/bin/env python3
"""Standalone referential-integrity checker for junctions_topology_v6.xlsx.

Independent of mapmaking.py's render pipeline on purpose: a rendering bug
only shows up if you happen to look at the right spot on the map (that's
how the Mosel/P070 mistake sat unnoticed - a tankstation wired into the
river chain instead of a river bridge). This script instead walks every
cross-sheet reference in the workbook and reports every broken one in one
pass, so a mistake is caught by running a check, not by spotting a glitch.

Usage: python3 check_data_integrity.py [path-to-xlsx]
Exits 1 if any error was found, 0 otherwise (warnings never affect the exit code).
"""
import math
import sys

from openpyxl import load_workbook

DEFAULT_XLSX = "junctions_topology_v6.xlsx"


def sheet_rows(wb, name):
    ws = wb[name]
    header = [c.value for c in ws[1]]
    idx = {h: i for i, h in enumerate(header) if h is not None}
    return idx, list(ws.iter_rows(min_row=2, values_only=True))


def check_duplicate_ids(wb, sheet, id_col, errors):
    idx, rows = sheet_rows(wb, sheet)
    seen = {}
    for row in rows:
        val = row[idx[id_col]]
        if val is None:
            continue
        seen.setdefault(val, 0)
        seen[val] += 1
    for val, count in seen.items():
        if count > 1:
            errors.append(f"[{sheet}] duplicate {id_col} \"{val}\" ({count} rows)")


def check_fk(wb, sheet, col, target_ids, target_label, errors, row_label_cols=()):
    idx, rows = sheet_rows(wb, sheet)
    for row in rows:
        val = row[idx[col]]
        if val is None or val in target_ids:
            continue
        label = " ".join(str(row[idx[c]]) for c in row_label_cols if c in idx and row[idx[c]] is not None)
        errors.append(f"[{sheet}] {col}=\"{val}\" ({label}) not found in {target_label}")


def check_position_range(wb, sheet, col, id_col, errors):
    idx, rows = sheet_rows(wb, sheet)
    for row in rows:
        val = row[idx[col]]
        if val is None:
            continue
        if not (0.0 <= val <= 1.0):
            errors.append(f"[{sheet}] {id_col}=\"{row[idx[id_col]]}\" has {col}={val}, outside [0, 1]")


def check_river_segments(wb, points_by_id, river_junction_ids, errors):
    idx, rows = sheet_rows(wb, "River Segments")
    for row in rows:
        fid, tid, river = row[idx["From ID"]], row[idx["To ID"]], row[idx.get("River")]
        if not fid or not tid:
            continue
        for node_id in (fid, tid):
            if node_id in river_junction_ids:
                continue
            pt = points_by_id.get(node_id)
            if pt is None:
                errors.append(f"[River Segments] {fid} -> {tid} ({river}): "
                              f"{node_id} is not a known Point ID or River Junction ID")
            elif pt["category"] != "brug (rivier)":
                errors.append(f"[River Segments] {fid} -> {tid} ({river}): "
                              f"{node_id} \"{pt['name']}\" has Category \"{pt['category']}\", "
                              f"not \"brug (rivier)\"")


def check_river_junction_degree(wb, warnings):
    """A confluence or split is where >=3 river stretches meet (two in, one
    out, or the reverse); anything else (border crossing, name-change,
    waypoint) is just one river passing through, so needs only 2. Fewer
    than that means some stretch that should reach this junction doesn't -
    exactly the RJ007/RJ008 gaps found by hand this session (a tributary's
    own inflow existed, but nothing carried the main river through)."""
    rjidx, rjrows = sheet_rows(wb, "River Junctions")
    rsidx, rsrows = sheet_rows(wb, "River Segments")
    degree = {}
    for row in rsrows:
        degree[row[rsidx["From ID"]]] = degree.get(row[rsidx["From ID"]], 0) + 1
        degree[row[rsidx["To ID"]]] = degree.get(row[rsidx["To ID"]], 0) + 1
    for row in rjrows:
        rid = row[rjidx["River Junction ID"]]
        if not rid:
            continue
        rtype = (row[rjidx["Type"]] or "").lower()
        expected = 3 if rtype in ("confluence", "split") else 2
        d = degree.get(rid, 0)
        if d < expected:
            warnings.append(f"[River Junctions] {rid} \"{row[rjidx['Name']]}\" (type {rtype or '?'}) "
                            f"has only {d} connected River Segment(s), expected >= {expected}")


def check_orphans(wb, warnings):
    """Nodes nothing ever connects to - not wrong, but worth a look."""
    jidx, jrows = sheet_rows(wb, "Junctions")
    sidx, srows = sheet_rows(wb, "Segments")
    referenced = set()
    for row in srows:
        referenced.add(row[sidx["From ID"]])
        referenced.add(row[sidx["To ID"]])
    for row in jrows:
        jid = row[jidx["ID"]]
        if jid and jid not in referenced:
            warnings.append(f"[Junctions] {jid} \"{row[jidx['Junction name']]}\" "
                            f"is never used as a Segments From/To ID")

    rjidx, rjrows = sheet_rows(wb, "River Junctions")
    rsidx, rsrows = sheet_rows(wb, "River Segments")
    river_referenced = set()
    for row in rsrows:
        river_referenced.add(row[rsidx["From ID"]])
        river_referenced.add(row[rsidx["To ID"]])
    for row in rjrows:
        rid = row[rjidx["River Junction ID"]]
        if rid and rid not in river_referenced:
            warnings.append(f"[River Junctions] {rid} \"{row[rjidx['Name']]}\" "
                            f"is never used as a River Segments From/To ID")

    bnidx, bnrows = sheet_rows(wb, "Border Nodes")
    bsidx, bsrows = sheet_rows(wb, "Border Segments")
    border_referenced = set()
    for row in bsrows:
        border_referenced.add(row[bsidx["FromBorderNodeID"]])
        border_referenced.add(row[bsidx["ToBorderNodeID"]])
    for row in bnrows:
        bnid = row[bnidx["BorderNodeID"]]
        if bnid and bnid not in border_referenced:
            warnings.append(f"[Border Nodes] {bnid} \"{row[bnidx['Name']]}\" "
                            f"is never used as a Border Segments From/To ID")


def check_ambiguous_one_sided_points(wb, warnings, near_km=0.3):
    """A one-sided station (Sides=1) with no 'Side direction' relies on its
    own coordinate to tell which side of the road it's on - a coordinate
    that was estimated by interpolating onto the straight line between the
    segment's two junctions (common for a station added from a brand
    survey rather than individually geocoded) sits so close to that line
    that the side becomes a coin flip, just as it did for Velder before it
    was fixed (see mapmaking.py's fuel_side_visible()/one_sided_marker_side()).
    Flags any one-sided, no-Side-direction point within `near_km` of its
    segment's own junction-to-junction line."""
    pidx, prows = sheet_rows(wb, "Points")
    jidx, jrows = sheet_rows(wb, "Junctions")
    sidx, srows = sheet_rows(wb, "Segments")
    junction_pos = {row[jidx["ID"]]: (row[jidx["Latitude"]], row[jidx["Longitude"]])
                    for row in jrows if row[jidx["ID"]] and row[jidx["Latitude"]] is not None}
    segment_ends = {row[sidx["Edge ID"]]: (row[sidx["From ID"]], row[sidx["To ID"]])
                    for row in srows if row[sidx["Edge ID"]]}
    for row in prows:
        if row[pidx["Sides (1 or 2)"]] != 1 or row[pidx.get("Side direction")]:
            continue
        lat, lon = row[pidx["Latitude"]], row[pidx["Longitude"]]
        eid = row[pidx["Edge ID"]]
        if lat is None or lon is None or eid not in segment_ends:
            continue
        fid, tid = segment_ends[eid]
        if fid not in junction_pos or tid not in junction_pos:
            continue
        flat, flon = junction_pos[fid]
        tlat, tlon = junction_pos[tid]
        cos_lat = math.cos(math.radians((flat + tlat) / 2))
        fx, fy = flon * cos_lat, flat
        tx, ty = tlon * cos_lat, tlat
        px, py = lon * cos_lat, lat
        road_dx, road_dy = tx - fx, ty - fy
        road_len = math.hypot(road_dx, road_dy)
        if road_len < 1e-9:
            continue
        cross = road_dx * (py - fy) - road_dy * (px - fx)
        dist_km = abs(cross) / road_len * 111.2
        if dist_km < near_km:
            warnings.append(f"[Points] {row[pidx['Point ID']]} \"{row[pidx['Name']]}\" is one-sided with no "
                            f"Side direction, and its coordinate is only ~{dist_km*1000:.0f}m from the "
                            f"{eid} From/To line - which side it renders on may be unreliable")


def main():
    xlsx_path = sys.argv[1] if len(sys.argv) > 1 else DEFAULT_XLSX
    wb = load_workbook(xlsx_path, data_only=True)
    errors = []
    warnings = []

    # 1. duplicate primary keys
    check_duplicate_ids(wb, "Junctions", "ID", errors)
    check_duplicate_ids(wb, "Segments", "Edge ID", errors)
    check_duplicate_ids(wb, "Points", "Point ID", errors)
    check_duplicate_ids(wb, "Rivers", "RiverID", errors)
    check_duplicate_ids(wb, "River Junctions", "River Junction ID", errors)
    check_duplicate_ids(wb, "Border Segments", "BorderSegmentID", errors)
    check_duplicate_ids(wb, "Border Nodes", "BorderNodeID", errors)

    # 2. foreign keys
    jidx, jrows = sheet_rows(wb, "Junctions")
    junction_ids = {row[jidx["ID"]] for row in jrows if row[jidx["ID"]]}
    check_fk(wb, "Segments", "From ID", junction_ids, "Junctions", errors, ["Edge ID", "From name"])
    check_fk(wb, "Segments", "To ID", junction_ids, "Junctions", errors, ["Edge ID", "To name"])

    sidx, srows = sheet_rows(wb, "Segments")
    edge_ids = {row[sidx["Edge ID"]] for row in srows if row[sidx["Edge ID"]]}
    check_fk(wb, "Points", "Edge ID", edge_ids, "Segments", errors, ["Point ID", "Name"])
    check_fk(wb, "Points of Interest", "Segment", edge_ids, "Segments", errors, ["Name"])

    check_position_range(wb, "Points", "PositionOnEdge", "Point ID", errors)

    pidx, prows = sheet_rows(wb, "Points")
    points_by_id = {row[pidx["Point ID"]]: {"name": row[pidx["Name"]], "category": row[pidx["Category"]]}
                    for row in prows if row[pidx["Point ID"]]}
    rjidx, rjrows = sheet_rows(wb, "River Junctions")
    river_junction_ids = {row[rjidx["River Junction ID"]] for row in rjrows if row[rjidx["River Junction ID"]]}
    check_river_segments(wb, points_by_id, river_junction_ids, errors)

    bnidx, bnrows = sheet_rows(wb, "Border Nodes")
    border_node_ids = {row[bnidx["BorderNodeID"]] for row in bnrows if row[bnidx["BorderNodeID"]]}
    check_fk(wb, "Border Segments", "FromBorderNodeID", border_node_ids, "Border Nodes", errors,
             ["BorderSegmentID", "FromName"])
    check_fk(wb, "Border Segments", "ToBorderNodeID", border_node_ids, "Border Nodes", errors,
             ["BorderSegmentID", "ToName"])
    check_fk(wb, "Border Nodes", "LinkedJunctionID", junction_ids, "Junctions", errors, ["BorderNodeID", "Name"])
    check_fk(wb, "Border Nodes", "LinkedRiverJunctionID", river_junction_ids, "River Junctions", errors,
             ["BorderNodeID", "Name"])

    nidx, nrows = sheet_rows(wb, "NaturalRegions")
    junction_names = {row[jidx["Junction name"]] for row in jrows if row[jidx["Junction name"]]}
    check_fk(wb, "NaturalRegions", "Junction name", junction_names, "Junctions", errors, ["Region name"])

    # 3. orphans and ambiguous one-sided points (soft)
    check_orphans(wb, warnings)
    check_ambiguous_one_sided_points(wb, warnings)
    check_river_junction_degree(wb, warnings)

    for e in errors:
        print("ERROR:", e)
    for w in warnings:
        print("warning:", w)
    print(f"\n{len(errors)} error(s), {len(warnings)} warning(s)")
    return 1 if errors else 0


if __name__ == "__main__":
    sys.exit(main())
